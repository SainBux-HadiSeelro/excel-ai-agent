"""
Script to generate sample data for Excel AI Agent testing and demonstrations.
Generates:
1. sample_data/sales.xlsx (with realistic sales, intentional duplicates, messy text, and formulas)
2. sample_data/regional_extra.xlsx (for testing merge_workbooks)
"""

import os
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

os.makedirs("sample_data", exist_ok=True)

wb = openpyxl.Workbook()
# Setup primary sheet
ws = wb.active
ws.title = "Sales"

headers = ["Order_ID", "Date", "Region", "Sales_Rep", "Product", "Units", "Unit_Price", "Total_Revenue", "Status"]
ws.append(headers)

base_date = date(2025, 1, 1)

rows = [
    ["ORD-1001", base_date + timedelta(days=1), "North", "Ali Khan", "Laptop", 5, 1200.0, "=F2*G2", "Completed"],
    ["ORD-1002", base_date + timedelta(days=2), "South", "Fatima Noor ", "Mouse", 25, 25.0, "=F3*G3", "Completed"],
    ["ORD-1003", base_date + timedelta(days=3), "East", "bilal ahmed", "Keyboard", 15, 45.0, "=F4*G4", "Pending"],
    ["ORD-1004", base_date + timedelta(days=4), "West", "Zainab Bibi", "Monitor", 8, 300.0, "=F5*G5", "Completed"],
    ["ORD-1005", base_date + timedelta(days=5), "North", "Ali Khan", "Laptop", 3, 1200.0, "=F6*G6", "Completed"],
    ["ORD-1006", base_date + timedelta(days=6), "South", "Usman Tariq", "USB Drive", 50, 15.0, "=F7*G7", "Cancelled"],
    ["ORD-1007", base_date + timedelta(days=7), "East", "Fatima Noor", "Desk Chair", 4, 180.0, "=F8*G8", "Completed"],
    ["ORD-1008", base_date + timedelta(days=8), "North", "hamza malik", "Webcam", 12, 60.0, "=F9*G9", "Completed"],
    ["ORD-1009", base_date + timedelta(days=9), "West", "Zainab Bibi", "Headphones", 20, 85.0, "=F10*G10", "Pending"],
    ["ORD-1010", base_date + timedelta(days=10), "South", "Usman Tariq", "Laptop", 6, 1250.0, "=F11*G11", "Completed"],
    # Intentional duplicate of ORD-1005 to test remove_duplicates:
    ["ORD-1005", base_date + timedelta(days=5), "North", "Ali Khan", "Laptop", 3, 1200.0, "=F12*G12", "Completed"],
    ["ORD-1011", base_date + timedelta(days=11), "East", "bilal ahmed", "Monitor", 7, 310.0, "=F13*G13", "Completed"],
    ["ORD-1012", base_date + timedelta(days=12), "North", "Ali Khan ", "Mouse", 30, 24.0, "=F14*G14", "Completed"],
    ["ORD-1013", base_date + timedelta(days=13), "West", "Sara Raza", "Keyboard", 18, 48.0, "=F15*G15", "Completed"],
    ["ORD-1014", base_date + timedelta(days=14), "South", "Fatima Noor", "Printer", 2, 450.0, "=F16*G16", "Pending"],
    ["ORD-1015", base_date + timedelta(days=15), "East", "Hamza Malik", "Laptop", 4, 1200.0, "=F17*G17", "Completed"],
    # Another duplicate row:
    ["ORD-1013", base_date + timedelta(days=13), "West", "Sara Raza", "Keyboard", 18, 48.0, "=F18*G18", "Completed"],
    ["ORD-1016", base_date + timedelta(days=16), "North", "Ali Khan", "Monitor", 10, 305.0, "=F19*G19", "Completed"],
    ["ORD-1017", base_date + timedelta(days=17), "South", "Usman Tariq", "Mouse", 40, 22.0, "=F20*G20", "Completed"],
    ["ORD-1018", base_date + timedelta(days=18), "West", "Sara Raza ", "Desk Chair", 5, 185.0, "=F21*G21", "Completed"],
    ["ORD-1019", base_date + timedelta(days=19), "East", "Bilal Ahmed", "USB Drive", 60, 14.0, "=F22*G22", "Completed"],
    ["ORD-1020", base_date + timedelta(days=20), "North", "Hamza Malik", "Headphones", 14, 80.0, "=F23*G23", "Completed"],
]

for r in rows:
    ws.append(r)

# Header styling
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Formatting data rows
for r_idx in range(2, len(rows) + 2):
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.border = thin_border
        if c_idx in (6,): # Units
            cell.number_format = "#,##0"
        elif c_idx in (7, 8): # Price, Revenue
            cell.number_format = "$#,##0.00"
        elif c_idx == 2: # Date
            cell.number_format = "YYYY-MM-DD"

# Auto-fit column widths
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Sheet 2: Regional Targets
ws_targets = wb.create_sheet(title="Targets")
target_headers = ["Region", "Target_Revenue", "Headcount"]
ws_targets.append(target_headers)
targets_data = [
    ["North", 25000, 5],
    ["South", 20000, 4],
    ["East", 18000, 3],
    ["West", 22000, 4],
]
for row in targets_data:
    ws_targets.append(row)

for col_num in range(1, len(target_headers) + 1):
    c = ws_targets.cell(row=1, column=col_num)
    c.font = header_font
    c.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

for col in ws_targets.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_targets.column_dimensions[col_letter].width = max(max_len + 3, 14)

wb.save("sample_data/sales.xlsx")

# Second workbook for testing merge
wb_extra = openpyxl.Workbook()
ws_extra = wb_extra.active
ws_extra.title = "Q1_Summary"
ws_extra.append(["Category", "Q1_Total", "Notes"])
ws_extra.append(["Hardware", 45000, "Strong demand"])
ws_extra.append(["Accessories", 12000, "Stock replenishment"])
ws_extra.append(["Furniture", 8500, "New branch openings"])
wb_extra.save("sample_data/regional_extra.xlsx")

print("Created sample_data/sales.xlsx and sample_data/regional_extra.xlsx successfully!")

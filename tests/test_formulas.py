"""Tests for advanced formula evaluation (SUMIF, AVERAGE, IF, and nested formulas)."""

import openpyxl
from workspace_manager import evaluate_formula_cell


def test_evaluate_formulas_advanced():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CalcSheet"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = 30

    ws["B1"] = "=AVERAGE(A1:A3)"
    ws["B2"] = '=IF(A1>5, "Pass", "Fail")'
    ws["B3"] = '=SUMIF(A1:A3, ">15")'
    ws["B4"] = '=IF(SUM(A1:A3)>50, AVERAGE(A1:A3)*2, 0)'

    # AVERAGE test
    val_avg = evaluate_formula_cell(ws["B1"].value, None, ws, cell_coord="B1")
    assert round(float(val_avg), 2) == 20.0

    # IF test
    val_if = evaluate_formula_cell(ws["B2"].value, None, ws, cell_coord="B2")
    assert str(val_if).strip() == "Pass"

    # SUMIF test
    val_sumif = evaluate_formula_cell(ws["B3"].value, None, ws, cell_coord="B3")
    assert round(float(val_sumif), 2) == 50.0

    # Nested formula test: IF + SUM + AVERAGE
    val_nested = evaluate_formula_cell(ws["B4"].value, None, ws, cell_coord="B4")
    assert round(float(val_nested), 2) == 40.0


def test_evaluate_formula_read_range_integration(sample_manager):
    from tools.editing import write_cells
    from tools.inspection import read_range

    # Write test formulas to Sales sheet
    write_cells(sample_manager, updates=[
        {"cell": "K2", "value": "=AVERAGE(F2:F5)"},
        {"cell": "K3", "value": '=IF(F2>10, "HighQty", "LowQty")'},
        {"cell": "K4", "value": '=SUMIF(F2:F5, ">5")'},
        {"cell": "K5", "value": '=IF(SUM(F2:F5)>0, AVERAGE(F2:F5), -1)'}
    ], sheet_name="Sales")

    # Read without show_formulas (should evaluate)
    res = read_range(sample_manager, sheet_name="Sales", cell_range="K2:K5", show_formulas=False)
    assert res["row_count"] == 4
    # Values should be evaluated, not raw formulas
    assert not str(res["rows"][0][0]).startswith("=")
    assert str(res["rows"][1][0]) in ("HighQty", "LowQty")
    assert float(res["rows"][2][0]) > 0
    assert float(res["rows"][3][0]) > 0


def test_formula_evaluation_failure_never_crashes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=COMPLETELY_INVALID_UNKNOWN_FUNC(999, 'bad')"
    # Should safely return raw formula without raising
    val = evaluate_formula_cell(ws["A1"].value, None, ws, cell_coord="A1")
    assert val == "=COMPLETELY_INVALID_UNKNOWN_FUNC(999, 'bad')"


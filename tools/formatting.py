"""
Formatting tools for Excel AI Agent:
- format_range (bold, colors, number formats, borders, alignments, column auto-fit)
"""

from __future__ import annotations

from typing import Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from workspace_manager import WorkspaceManager


def format_range(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    cell_range: Optional[str] = None,
    bold: Optional[bool] = None,
    italic: Optional[bool] = None,
    font_color: Optional[str] = None,
    bg_color: Optional[str] = None,
    font_size: Optional[int] = None,
    font_name: Optional[str] = None,
    number_format: Optional[str] = None,
    border: Optional[str] = None,  # 'thin', 'medium', 'double', 'thick'
    border_color: str = "D9D9D9",
    align_horizontal: Optional[str] = None,  # 'left', 'center', 'right'
    align_vertical: Optional[str] = None,  # 'top', 'center', 'bottom'
    auto_fit_columns: bool = False
) -> Dict[str, Any]:
    """
    Applies rich visual formatting to a range or sheet:
    - font style (bold, italic, size, color, name)
    - cell background fill color
    - number formatting (currency, percent, dates, integers)
    - borders
    - horizontal/vertical text alignment
    - auto-fit column dimensions based on text length
    """
    wb = manager.open_workbook(data_only=False)

    # Clean sheet prefix if present
    if cell_range and "!" in cell_range:
        s_part, _, r_part = cell_range.rpartition("!")
        clean_s = s_part.strip().strip("'")
        if not sheet_name and clean_s in wb.sheetnames:
            sheet_name = clean_s
        cell_range = r_part.strip().strip("'")

    if sheet_name and sheet_name not in wb.sheetnames:
        return {"error": f"Sheet '{sheet_name}' does not exist. Available: {wb.sheetnames}"}
    ws = wb[sheet_name] if sheet_name else wb.active

    formatted_count = 0

    # Auto-fit columns if requested
    if auto_fit_columns:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    if cell_range:
        try:
            cells = ws[cell_range]
            # Normalize to 2D list
            if not isinstance(cells, tuple):
                cell_matrix = [[cells]]
            elif len(cells) > 0 and not isinstance(cells[0], tuple):
                cell_matrix = [cells]
            else:
                cell_matrix = cells
        except Exception as e:
            return {"error": f"Invalid cell range '{cell_range}': {str(e)}"}

        clean_bg = bg_color.lstrip("#") if bg_color else None
        clean_fg = font_color.lstrip("#") if font_color else None
        clean_border = border_color.lstrip("#") if border_color else "D9D9D9"

        border_side = Side(style=border, color=clean_border) if border else None
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side) if border_side else None

        for row in cell_matrix:
            for cell in row:
                # Font updates
                if any(x is not None for x in (bold, italic, font_size, font_name, clean_fg)):
                    current_font = cell.font or Font()
                    new_bold = bold if bold is not None else current_font.bold
                    new_italic = italic if italic is not None else current_font.italic
                    new_size = font_size if font_size is not None else current_font.size
                    new_name = font_name if font_name is not None else current_font.name
                    new_color = clean_fg if clean_fg is not None else (current_font.color if current_font else None)

                    cell.font = Font(
                        name=new_name,
                        size=new_size,
                        bold=new_bold,
                        italic=new_italic,
                        color=new_color
                    )

                # Fill update
                if clean_bg:
                    cell.fill = PatternFill(start_color=clean_bg, end_color=clean_bg, fill_type="solid")

                # Number format
                if number_format:
                    cell.number_format = number_format

                # Alignment
                current_align = cell.alignment or Alignment()
                new_h = align_horizontal if align_horizontal is not None else current_align.horizontal
                new_v = align_vertical if align_vertical is not None else current_align.vertical
                if align_horizontal or align_vertical:
                    cell.alignment = Alignment(horizontal=new_h, vertical=new_v)

                # Border
                if cell_border:
                    cell.border = cell_border

                formatted_count += 1

    manager.prepare_for_mutation(f"format_range: {cell_range or 'sheet'}")
    manager.save_workbook(wb)

    return {
        "success": True,
        "sheet_name": ws.title,
        "cell_range": cell_range or "auto-fit only",
        "formatted_cells_count": formatted_count,
        "auto_fit_applied": auto_fit_columns,
        "message": f"Successfully applied formatting to {formatted_count} cell(s) in '{ws.title}'."
    }

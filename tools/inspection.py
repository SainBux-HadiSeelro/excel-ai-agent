"""
Inspection tools for Excel AI Agent:
- list_sheets
- read_range
- get_sheet_summary
"""

from __future__ import annotations

from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.utils import get_column_letter
from workspace_manager import WorkspaceManager, evaluate_formula_cell


def list_sheets(manager: WorkspaceManager) -> Dict[str, Any]:
    """Lists all sheet names and properties in the current workbook."""
    wb = manager.open_workbook()
    sheets = []
    for idx, name in enumerate(wb.sheetnames):
        ws = wb[name]
        sheets.append({
            "index": idx,
            "name": name,
            "max_row": ws.max_row or 0,
            "max_column": ws.max_column or 0,
            "is_active": (name == wb.active.title if wb.active else False)
        })
    return {
        "sheets": sheets,
        "total_sheets": len(sheets),
        "active_sheet": wb.active.title if wb.active else None
    }


def read_range(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    cell_range: Optional[str] = None,
    max_rows: int = 50,
    show_formulas: bool = True
) -> Dict[str, Any]:
    """
    Reads a specified range (e.g. 'A1:D20' or entire sheet) with formula visibility and row limiting.
    """
    wb = manager.open_workbook(data_only=not show_formulas)
    wb_raw = manager.open_workbook(data_only=False) if not show_formulas else None

    if cell_range and "!" in cell_range:
        sheet_part, _, range_part = cell_range.rpartition("!")
        clean_sheet = sheet_part.strip().strip("'")
        cell_range = range_part.strip().strip("'")
        if not sheet_name and clean_sheet in wb.sheetnames:
            sheet_name = clean_sheet

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            return {"error": f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}"}
        ws = wb[sheet_name]
    else:
        ws = wb.active

    raw_ws = wb_raw[ws.title] if wb_raw else None

    if ws.max_row == 0 or ws.max_column == 0:
        return {
            "sheet_name": ws.title,
            "cell_range": cell_range or "Empty",
            "rows": [],
            "row_count": 0,
            "message": "Sheet is empty."
        }

    extracted_rows: List[List[Any]] = []

    if cell_range:
        try:
            cells = ws[cell_range]
            # If single cell
            if not isinstance(cells, tuple):
                if not show_formulas and raw_ws:
                    raw_cell = raw_ws[cells.coordinate]
                    val = evaluate_formula_cell(raw_cell.value, cells.value, raw_ws, ws, cell_coord=cells.coordinate)
                else:
                    val = cells.value if cells.value is not None else ""
                return {
                    "sheet_name": ws.title,
                    "cell_range": cell_range,
                    "rows": [[val]],
                    "row_count": 1,
                    "truncated": False
                }
            for row in cells[:max_rows]:
                row_vals = []
                for c in row:
                    if not show_formulas and raw_ws:
                        raw_c = raw_ws.cell(row=c.row, column=c.column)
                        row_vals.append(evaluate_formula_cell(raw_c.value, c.value, raw_ws, ws, cell_coord=raw_c.coordinate))
                    else:
                        row_vals.append(c.value if c.value is not None else "")
                extracted_rows.append(row_vals)
            truncated = len(cells) > max_rows
        except Exception as e:
            return {"error": f"Invalid range expression '{cell_range}': {str(e)}"}
    else:
        # Default: read from top up to max_rows
        limit_rows = min(ws.max_row, max_rows)
        limit_cols = ws.max_column
        for r in range(1, limit_rows + 1):
            row_data = []
            for c in range(1, limit_cols + 1):
                val = ws.cell(row=r, column=c).value
                if not show_formulas and raw_ws:
                    raw_val = raw_ws.cell(row=r, column=c).value
                    coord = f"{get_column_letter(c)}{r}"
                    row_data.append(evaluate_formula_cell(raw_val, val, raw_ws, ws, cell_coord=coord))
                else:
                    row_data.append(val if val is not None else "")
            extracted_rows.append(row_data)
        truncated = ws.max_row > max_rows

    return {
        "sheet_name": ws.title,
        "cell_range": cell_range or f"A1:{get_column_letter(ws.max_column)}{min(ws.max_row, max_rows)}",
        "rows": extracted_rows,
        "row_count": len(extracted_rows),
        "total_sheet_rows": ws.max_row,
        "total_sheet_cols": ws.max_column,
        "truncated": truncated
    }


def get_sheet_summary(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    max_sample_rows: int = 5
) -> Dict[str, Any]:
    """
    Summarizes sheet: column headers, inferred types, null counts, row counts, and sample rows.
    Handles large files (50k+ rows) gracefully without overwhelming memory.
    """
    wb = manager.open_workbook(data_only=False)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            return {"error": f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"}
        ws = wb[sheet_name]
    else:
        ws = wb.active

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    if max_row == 0 or max_col == 0:
        return {
            "sheet_name": ws.title,
            "row_count": 0,
            "col_count": 0,
            "columns": [],
            "message": "Sheet is empty."
        }

    # Extract headers from row 1
    headers = []
    for c in range(1, max_col + 1):
        val = ws.cell(row=1, column=c).value
        headers.append(str(val) if val is not None else f"Column_{get_column_letter(c)}")

    # Sample rows
    sample_rows = []
    sample_limit = min(max_row, max_sample_rows + 1)
    for r in range(2, sample_limit + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        sample_rows.append(row_vals)

    # Lightweight statistics scan (up to first 5,000 rows to keep it ultra fast even on 100k files)
    scan_limit = min(max_row, 5000)
    col_nulls = [0] * max_col
    col_types = [set() for _ in range(max_col)]

    for r in range(2, scan_limit + 1):
        for c_idx in range(max_col):
            v = ws.cell(row=r, column=c_idx + 1).value
            if v is None or v == "":
                col_nulls[c_idx] += 1
            else:
                if isinstance(v, str) and v.startswith("="):
                    col_types[c_idx].add("formula")
                else:
                    col_types[c_idx].add(type(v).__name__)

    columns_info = []
    for i, h in enumerate(headers):
        types_list = sorted(list(col_types[i]))
        inferred_type = "/".join(types_list) if types_list else "empty"
        columns_info.append({
            "column_index": i + 1,
            "column_letter": get_column_letter(i + 1),
            "name": h,
            "inferred_type": inferred_type,
            "null_count_sample": col_nulls[i]
        })

    return {
        "sheet_name": ws.title,
        "total_rows": max_row,
        "total_columns": max_col,
        "scanned_rows": scan_limit,
        "columns": columns_info,
        "sample_rows": sample_rows
    }

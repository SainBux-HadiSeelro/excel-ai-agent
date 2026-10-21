"""
Sheet and structure manipulation tools:
- create_sheet
- delete_sheet (destructive guard)
- rename_sheet
- insert_rows
- delete_rows (destructive guard)
- insert_columns
- delete_columns (destructive guard)
"""

from __future__ import annotations

from typing import Dict, Any, Optional
from openpyxl.utils import get_column_letter, column_index_from_string
from workspace_manager import WorkspaceManager


def create_sheet(
    manager: WorkspaceManager,
    title: str,
    index: Optional[int] = None
) -> Dict[str, Any]:
    """Creates a new empty sheet with the given title."""
    if not title or not title.strip():
        return {"error": "Sheet title cannot be empty."}

    title = title.strip()
    wb = manager.open_workbook()

    if title in wb.sheetnames:
        return {"error": f"Sheet with name '{title}' already exists."}

    manager.prepare_for_mutation(f"create_sheet: {title}")
    ws = wb.create_sheet(title=title, index=index)
    manager.save_workbook(wb)

    return {
        "success": True,
        "created_sheet": title,
        "index": wb.sheetnames.index(title),
        "total_sheets": len(wb.sheetnames),
        "message": f"Successfully created sheet '{title}'."
    }


def delete_sheet(
    manager: WorkspaceManager,
    title: str,
    confirmed: bool = False
) -> Dict[str, Any]:
    """Deletes an existing sheet. Destructive operation requiring confirmation."""
    wb = manager.open_workbook()

    if title not in wb.sheetnames:
        return {"error": f"Sheet '{title}' does not exist in workbook. Available: {wb.sheetnames}"}

    if len(wb.sheetnames) <= 1:
        return {"error": "Cannot delete the only sheet in a workbook."}

    target_sheet = wb[title]
    row_count = target_sheet.max_row or 0
    col_count = target_sheet.max_column or 0

    if not confirmed:
        return {
            "requires_confirmation": True,
            "destructive_action": "delete_sheet",
            "sheet_name": title,
            "warning": f"Sheet '{title}' contains {row_count} rows and {col_count} columns. All data will be deleted.",
            "message": f"Confirmation required to delete sheet '{title}'."
        }

    manager.prepare_for_mutation(f"delete_sheet: {title}")
    wb.remove(target_sheet)
    manager.save_workbook(wb)

    return {
        "success": True,
        "deleted_sheet": title,
        "remaining_sheets": wb.sheetnames,
        "message": f"Successfully deleted sheet '{title}'."
    }


def rename_sheet(
    manager: WorkspaceManager,
    old_title: str,
    new_title: str
) -> Dict[str, Any]:
    """Renames an existing sheet."""
    if not new_title or not new_title.strip():
        return {"error": "New sheet title cannot be empty."}

    old_title = old_title.strip()
    new_title = new_title.strip()

    wb = manager.open_workbook()

    if old_title not in wb.sheetnames:
        return {"error": f"Sheet '{old_title}' does not exist. Available: {wb.sheetnames}"}

    if new_title in wb.sheetnames and new_title.lower() != old_title.lower():
        return {"error": f"A sheet named '{new_title}' already exists."}

    manager.prepare_for_mutation(f"rename_sheet: {old_title} -> {new_title}")
    ws = wb[old_title]
    ws.title = new_title
    manager.save_workbook(wb)

    return {
        "success": True,
        "old_title": old_title,
        "new_title": new_title,
        "all_sheets": wb.sheetnames,
        "message": f"Renamed sheet '{old_title}' to '{new_title}'."
    }


def insert_rows(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    row_idx: int = 1,
    count: int = 1
) -> Dict[str, Any]:
    """Inserts one or more blank rows before row_idx."""
    if row_idx < 1:
        return {"error": "row_idx must be >= 1."}
    if count < 1:
        return {"error": "count must be >= 1."}

    wb = manager.open_workbook()
    if sheet_name and sheet_name not in wb.sheetnames:
        return {"error": f"Sheet '{sheet_name}' does not exist. Available: {wb.sheetnames}"}
    ws = wb[sheet_name] if sheet_name else wb.active

    manager.prepare_for_mutation(f"insert_rows at {row_idx} (count={count})")
    ws.insert_rows(idx=row_idx, amount=count)
    manager.save_workbook(wb)

    return {
        "success": True,
        "sheet_name": ws.title,
        "inserted_at_row": row_idx,
        "count": count,
        "new_max_row": ws.max_row,
        "message": f"Inserted {count} row(s) before row {row_idx} in sheet '{ws.title}'."
    }


def delete_rows(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    row_idx: int = 1,
    count: int = 1,
    confirmed: bool = False
) -> Dict[str, Any]:
    """Deletes one or more rows. Destructive operation requiring confirmation."""
    if row_idx < 1:
        return {"error": "row_idx must be >= 1."}
    if count < 1:
        return {"error": "count must be >= 1."}

    wb = manager.open_workbook()
    if sheet_name and sheet_name not in wb.sheetnames:
        return {"error": f"Sheet '{sheet_name}' does not exist. Available: {wb.sheetnames}"}
    ws = wb[sheet_name] if sheet_name else wb.active

    if not confirmed:
        # Sample preview of what's being deleted
        sample_preview = []
        end_idx = min(ws.max_row or 1, row_idx + count)
        for r in range(row_idx, end_idx):
            vals = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column + 1, 6))]
            sample_preview.append(vals)

        return {
            "requires_confirmation": True,
            "destructive_action": "delete_rows",
            "sheet_name": ws.title,
            "row_idx": row_idx,
            "count": count,
            "warning": f"About to delete {count} row(s) starting at row {row_idx}.",
            "sample_rows_to_delete": sample_preview,
            "message": f"Confirmation required to delete {count} row(s) starting from row {row_idx}."
        }

    manager.prepare_for_mutation(f"delete_rows at {row_idx} (count={count})")
    ws.delete_rows(idx=row_idx, amount=count)
    manager.save_workbook(wb)

    return {
        "success": True,
        "sheet_name": ws.title,
        "deleted_row_start": row_idx,
        "deleted_count": count,
        "new_max_row": ws.max_row,
        "message": f"Deleted {count} row(s) starting at row {row_idx} in sheet '{ws.title}'."
    }


def insert_columns(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    col_idx: int | str = 1,
    count: int = 1
) -> Dict[str, Any]:
    """Inserts one or more blank columns before col_idx (can be integer or column letter e.g. 'C')."""
    if isinstance(col_idx, str):
        col_number = column_index_from_string(col_idx.strip().upper())
    else:
        col_number = int(col_idx)

    if col_number < 1:
        return {"error": "col_idx must be >= 1."}
    if count < 1:
        return {"error": "count must be >= 1."}

    wb = manager.open_workbook()
    if sheet_name and sheet_name not in wb.sheetnames:
        return {"error": f"Sheet '{sheet_name}' does not exist. Available: {wb.sheetnames}"}
    ws = wb[sheet_name] if sheet_name else wb.active

    manager.prepare_for_mutation(f"insert_columns at {get_column_letter(col_number)} (count={count})")
    ws.insert_cols(idx=col_number, amount=count)
    manager.save_workbook(wb)

    return {
        "success": True,
        "sheet_name": ws.title,
        "inserted_at_column": get_column_letter(col_number),
        "count": count,
        "new_max_column": ws.max_column,
        "message": f"Inserted {count} column(s) before column {get_column_letter(col_number)} in '{ws.title}'."
    }


def delete_columns(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    col_idx: int | str = 1,
    count: int = 1,
    confirmed: bool = False
) -> Dict[str, Any]:
    """Deletes one or more columns. Destructive operation requiring confirmation."""
    if isinstance(col_idx, str):
        col_number = column_index_from_string(col_idx.strip().upper())
    else:
        col_number = int(col_idx)

    if col_number < 1:
        return {"error": "col_idx must be >= 1."}
    if count < 1:
        return {"error": "count must be >= 1."}

    wb = manager.open_workbook()
    if sheet_name and sheet_name not in wb.sheetnames:
        return {"error": f"Sheet '{sheet_name}' does not exist. Available: {wb.sheetnames}"}
    ws = wb[sheet_name] if sheet_name else wb.active

    col_letter = get_column_letter(col_number)

    if not confirmed:
        # Check header of the column
        header_val = ws.cell(row=1, column=col_number).value
        return {
            "requires_confirmation": True,
            "destructive_action": "delete_columns",
            "sheet_name": ws.title,
            "col_idx": col_number,
            "col_letter": col_letter,
            "count": count,
            "header_name": header_val,
            "warning": f"About to delete {count} column(s) starting at column {col_letter} ('{header_val}').",
            "message": f"Confirmation required to delete column(s) starting at {col_letter}."
        }

    manager.prepare_for_mutation(f"delete_columns at {col_letter} (count={count})")
    ws.delete_cols(idx=col_number, amount=count)
    manager.save_workbook(wb)

    return {
        "success": True,
        "sheet_name": ws.title,
        "deleted_column": col_letter,
        "deleted_count": count,
        "new_max_column": ws.max_column,
        "message": f"Deleted {count} column(s) starting at {col_letter} in sheet '{ws.title}'."
    }

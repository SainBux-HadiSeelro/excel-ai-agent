"""
Visualization tools for Excel AI Agent:
- add_chart (bar, line, pie)
"""

from __future__ import annotations

from typing import Dict, Any, Optional
import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils.cell import range_boundaries
from workspace_manager import WorkspaceManager


def add_chart(
    manager: WorkspaceManager,
    sheet_name: Optional[str] = None,
    chart_type: str = "bar",  # 'bar', 'column', 'line', 'pie'
    data_range: str = "",
    categories_range: Optional[str] = None,
    title: Optional[str] = None,
    target_cell: str = "K2",
    x_title: Optional[str] = None,
    y_title: Optional[str] = None
) -> Dict[str, Any]:
    """
    Creates and embeds a native Excel chart (Bar, Column, Line, or Pie) in the worksheet.
    - data_range: cell range containing numbers e.g. 'H1:H20' or 'B1:C10' (including header)
    - categories_range: cell range for X-axis labels e.g. 'C2:C20' (without header)
    - target_cell: where to anchor top-left corner of chart e.g. 'K2'
    """
    if not data_range:
        return {"error": "data_range must be provided (e.g. 'F1:F10')."}

    wb = manager.open_workbook(data_only=False)

    # Clean sheet prefix if present
    if data_range and "!" in data_range:
        s_part, _, r_part = data_range.rpartition("!")
        clean_s = s_part.strip().strip("'")
        if not sheet_name and clean_s in wb.sheetnames:
            sheet_name = clean_s
        data_range = r_part.strip().strip("'")

    if categories_range and "!" in categories_range:
        _, _, r_part = categories_range.rpartition("!")
        categories_range = r_part.strip().strip("'")

    if sheet_name and sheet_name not in wb.sheetnames:
        return {"error": f"Sheet '{sheet_name}' does not exist. Available: {wb.sheetnames}"}
    ws = wb[sheet_name] if sheet_name else wb.active

    ctype = chart_type.lower().strip()
    if ctype in ("bar", "bar_chart"):
        chart = BarChart()
        chart.type = "bar"
        chart_title = title or "Bar Chart"
    elif ctype in ("col", "column", "column_chart"):
        chart = BarChart()
        chart.type = "col"
        chart_title = title or "Column Chart"
    elif ctype in ("line", "line_chart"):
        chart = LineChart()
        chart_title = title or "Line Chart"
    elif ctype in ("pie", "pie_chart"):
        chart = PieChart()
        chart_title = title or "Pie Chart"
    else:
        return {"error": f"Unsupported chart type '{chart_type}'. Supported: 'bar', 'column', 'line', 'pie'."}

    chart.title = chart_title

    try:
        min_col, min_row, max_col, max_row = range_boundaries(data_range)
        data_ref = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        # Auto-detect whether data_range contains headers based on category row count
        titles_from_data = True
        if categories_range:
            c_min_col, c_min_row, c_max_col, c_max_row = range_boundaries(categories_range)
            cat_rows = c_max_row - c_min_row + 1
            data_rows = max_row - min_row + 1
            titles_from_data = (data_rows > cat_rows)
            cats_ref = Reference(ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
            chart.set_categories(cats_ref)

        chart.add_data(data_ref, titles_from_data=titles_from_data)

        if x_title and hasattr(chart, "x_axis"):
            chart.x_axis.title = x_title
        if y_title and hasattr(chart, "y_axis"):
            chart.y_axis.title = y_title

        chart.width = 15
        chart.height = 10
        manager.prepare_for_mutation(f"add_chart: {chart_type}")
        ws.add_chart(chart, target_cell)
        manager.save_workbook(wb)

        return {
            "success": True,
            "sheet_name": ws.title,
            "chart_type": ctype,
            "title": chart_title,
            "data_range": data_range,
            "target_cell": target_cell,
            "message": f"Successfully added {ctype} chart titled '{chart_title}' at cell {target_cell} in sheet '{ws.title}'."
        }
    except Exception as e:
        return {"error": f"Failed creating chart: {str(e)}"}

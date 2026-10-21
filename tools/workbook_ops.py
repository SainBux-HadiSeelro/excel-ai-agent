"""
Workbook-level operations for Excel AI Agent:
- merge_workbooks
- undo_last_change
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl
import pandas as pd
from workspace_manager import WorkspaceManager


def merge_workbooks(
    manager: WorkspaceManager,
    source_files: List[str],
    mode: str = "add_sheets",  # 'add_sheets' or 'append_rows'
    target_sheet_names: Optional[List[str]] = None
) -> Dict[str, Any]:
    """
    Merges external workbooks into the current workbook.
    - 'add_sheets': copies sheets from external workbooks into new sheets in current workbook.
    - 'append_rows': appends data rows from source sheets into the active sheet.
    """
    if not source_files:
        return {"error": "source_files list cannot be empty."}

    # Validate all sources before snapshotting or mutating anything
    for f_path in source_files:
        if not Path(f_path).exists():
            return {"error": f"Source file does not exist: {f_path}"}

    manager.prepare_for_mutation(f"merge_workbooks ({len(source_files)} files)")
    wb = manager.open_workbook()

    imported_sheets = []
    appended_rows_total = 0

    for idx, f_path in enumerate(source_files):
        p = Path(f_path)

        if p.suffix.lower() == ".csv":
            df = pd.read_csv(p)
            sheet_title = (
                target_sheet_names[idx]
                if target_sheet_names and idx < len(target_sheet_names)
                else p.stem
            )
            # Ensure unique sheet title
            base_title = sheet_title
            counter = 1
            while sheet_title in wb.sheetnames:
                sheet_title = f"{base_title}_{counter}"
                counter += 1

            if mode == "add_sheets":
                ws_new = wb.create_sheet(title=sheet_title)
                ws_new.append(list(df.columns))
                for row in df.itertuples(index=False):
                    ws_new.append(list(row))
                imported_sheets.append(sheet_title)
            else:
                ws_active = wb.active
                for row in df.itertuples(index=False):
                    ws_active.append(list(row))
                    appended_rows_total += 1
        else:
            # Excel file
            src_wb = openpyxl.load_workbook(p, data_only=False)
            for src_sheet_name in src_wb.sheetnames:
                src_ws = src_wb[src_sheet_name]

                if mode == "add_sheets":
                    new_title = (
                        target_sheet_names[idx]
                        if target_sheet_names and idx < len(target_sheet_names)
                        else f"{p.stem}_{src_sheet_name}"
                    )
                    # Deduplicate name
                    base_title = new_title
                    counter = 1
                    while new_title in wb.sheetnames:
                        new_title = f"{base_title}_{counter}"
                        counter += 1

                    new_ws = wb.create_sheet(title=new_title)
                    for row in src_ws.iter_rows(values_only=False):
                        new_row_vals = [c.value for c in row]
                        new_ws.append(new_row_vals)
                    imported_sheets.append(new_title)
                else:
                    ws_active = wb.active
                    # Append rows (skip header if active sheet already has headers)
                    start_r = 2 if ws_active.max_row > 1 else 1
                    for r_idx, row in enumerate(src_ws.iter_rows(values_only=False), start=1):
                        if r_idx < start_r:
                            continue
                        ws_active.append([c.value for c in row])
                        appended_rows_total += 1

    manager.save_workbook(wb)

    if mode == "add_sheets":
        return {
            "success": True,
            "mode": "add_sheets",
            "imported_sheets": imported_sheets,
            "all_sheets": wb.sheetnames,
            "message": f"Successfully imported {len(imported_sheets)} sheet(s): {imported_sheets}."
        }
    else:
        return {
            "success": True,
            "mode": "append_rows",
            "appended_rows_total": appended_rows_total,
            "active_sheet": wb.active.title,
            "message": f"Successfully appended {appended_rows_total} row(s) to '{wb.active.title}'."
        }


def undo_last_change(manager: WorkspaceManager) -> Dict[str, Any]:
    """Reverts the last modifying action using the snapshot history."""
    result = manager.undo()
    return result
